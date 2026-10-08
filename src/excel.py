import re

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.modelos import Asignatura, Clase


# --------------------------------------------------
# PATRONES
# --------------------------------------------------

# Una clase: abreviatura (sin espacios), tipo entre paréntesis
# y, opcionalmente, el aula entre < >.
PATRON_CLASE = re.compile(r"([^\s()<>]+)\s*\(([^)]+)\)(?:\s*<([^>]+)>)?")

PATRON_SEMANA = re.compile(r"Astea/Semana:\s*(\d+)", re.IGNORECASE)

PATRON_CURSO = re.compile(r"Kurtsoa/Curso:\s*(\d+)", re.IGNORECASE)

PATRON_GRUPO = re.compile(r"Taldea/Grupo:\s*(\d+)", re.IGNORECASE)

PATRON_HORA = re.compile(r"\d{1,2}:\d{2}")

INICIO_TABLA_ASIGNATURAS = "ASIGNATURA/IRAKASGAIA"

DIAS = {
    "lunes": "Lunes",
    "martes": "Martes",
    "miércoles": "Miércoles",
    "jueves": "Jueves",
    "viernes": "Viernes",
}


# --------------------------------------------------
# PUNTO DE ENTRADA
# --------------------------------------------------

def leer_horario(ruta) -> list[Clase]:
    """
    Lee todas las hojas del Excel y devuelve
    todas las clases encontradas.
    """

    libro = load_workbook(ruta, data_only=True)

    try:
        return [
            clase
            for hoja in libro.worksheets
            for clase in leer_hoja(hoja)
        ]

    finally:
        libro.close()


# --------------------------------------------------
# LECTURA DE UNA HOJA
# --------------------------------------------------

def leer_hoja(hoja: Worksheet) -> list[Clase]:
    """
    Lee todos los bloques semanales de una hoja.

    Cada bloque tiene esta forma:

        Kurtsoa/Curso ... Astea/Semana: X    <- inicio del bloque
        Lunes | Martes | ...                 <- fila de días
        8:40-9:30 | EM.II (T) <1.9> | ...    <- franjas horarias
        ...

    El bloque termina cuando empieza otra semana
    o la tabla inferior de asignaturas.
    """

    asignaturas = leer_asignaturas(hoja)

    clases = []

    semana = None
    curso = grupo = None
    columnas_dias = {}

    filas = hoja.iter_rows()

    for fila in filas:

        numero_semana = obtener_semana(fila)

        if numero_semana is not None:
            semana = numero_semana
            curso, grupo = obtener_curso_grupo(fila)

            # La fila siguiente contiene los días
            columnas_dias = obtener_columnas_dias(
                next(filas, ()),
                hoja.title,
                fila[0].row + 1,
            )
            continue

        primera_celda = fila[0].value

        if es_inicio_tabla_asignaturas(primera_celda):
            semana = None

        elif semana is not None and es_franja_horaria(primera_celda):
            clases.extend(
                leer_franja(fila, columnas_dias, semana, asignaturas, curso, grupo)
            )

    return clases


def leer_franja(fila, columnas_dias, semana, asignaturas, curso=None, grupo=None) -> list[Clase]:
    """
    Lee las clases de una fila horaria (una celda por día).
    """

    franja = str(fila[0].value)

    clases = []

    for dia, columna in columnas_dias.items():

        for abreviatura, tipo, aula in interpretar_celda(fila[columna].value):

            if abreviatura not in asignaturas:
                raise ValueError(
                    f"La asignatura '{abreviatura}' "
                    f"no aparece en la tabla de asignaturas."
                )

            clases.append(
                Clase(
                    asignatura=asignaturas[abreviatura],
                    tipo=tipo,
                    aula=aula,
                    semana=semana,
                    dia=dia,
                    franja=franja,
                    curso=curso,
                    grupo=grupo,
                )
            )

    return clases


# --------------------------------------------------
# CELDAS
# --------------------------------------------------

def interpretar_celda(valor):
    """
    Separa las clases que hay en una celda.

    Ejemplos:

        EM.II (T) <1.9>          ->  [("EM.II", "T", "1.9")]
        ELEK.D (GL1)             ->  [("ELEK.D", "GL1", None)]

    Una celda puede tener varias clases a la vez
    (dos laboratorios a la misma hora):

        MICR (GL4) <0.17>
        REGU (GL2)               ->  [("MICR", "GL4", "0.17"),
                                      ("REGU", "GL2", None)]

    Devuelve una lista vacía si la celda no contiene ninguna clase.
    """

    if valor is None:
        return []

    clases = []

    for abreviatura, tipo, aula in PATRON_CLASE.findall(str(valor)):

        # findall devuelve "" cuando no hay aula
        if aula == "":
            aula = None
        else:
            aula = aula.strip()

        clases.append((abreviatura, tipo.strip(), aula))

    return clases


def obtener_semana(fila):
    """
    Busca el número de semana en una fila.

    Ejemplo:

        Kurtsoa/Curso: 3 ... Astea/Semana: 4  ->  4
    """

    for celda in fila:

        if celda.value is None:
            continue

        coincidencia = PATRON_SEMANA.search(str(celda.value))

        if coincidencia:
            return int(coincidencia.group(1))

    return None


def obtener_curso_grupo(fila):
    """
    Busca el curso y el grupo en la cabecera de la semana.

    Ejemplo:

        Kurtsoa/Curso: 3  Taldea/Grupo: 1  ...  ->  ("3", "01")

    El grupo se devuelve con dos cifras, como en GAUR.
    """

    curso = grupo = None

    for celda in fila:

        if celda.value is None:
            continue

        texto = str(celda.value)

        if curso is None and (coincidencia := PATRON_CURSO.search(texto)):
            curso = coincidencia.group(1)

        if grupo is None and (coincidencia := PATRON_GRUPO.search(texto)):
            grupo = coincidencia.group(1).zfill(2)

    return curso, grupo


def obtener_columnas_dias(fila, nombre_hoja, numero_fila):
    """
    Detecta en qué columna está cada día.

    No importa que existan columnas vacías entre días.
    Devuelve la posición de cada día dentro de la fila.
    """

    columnas = {}

    for indice, celda in enumerate(fila):

        if celda.value is None:
            continue

        texto = str(celda.value).strip().lower()

        for palabra, dia in DIAS.items():

            if palabra in texto:
                columnas[dia] = indice
                break

    if len(columnas) != len(DIAS):
        raise ValueError(
            f"No se han encontrado los cinco días "
            f"en la fila {numero_fila} "
            f"de la hoja '{nombre_hoja}'."
        )

    return columnas


def es_franja_horaria(valor):
    """
    Comprueba si una celda contiene una franja horaria.

    Ejemplo:

        8:40-9:30
    """

    if valor is None:
        return False

    return len(PATRON_HORA.findall(str(valor))) >= 2


def es_inicio_tabla_asignaturas(valor):
    return valor is not None and INICIO_TABLA_ASIGNATURAS in str(valor)


# --------------------------------------------------
# ASIGNATURAS
# --------------------------------------------------

def leer_asignaturas(hoja: Worksheet) -> dict[str, Asignatura]:
    """
    Busca las filas de la tabla inferior de asignaturas.

    Ejemplo:

        EM.II | 26643 | Elektromagnetismoa II

    Devuelve un diccionario indexado por abreviatura.
    """

    asignaturas = {}

    for fila in hoja.iter_rows(values_only=True):

        if len(fila) < 3:
            continue

        abreviatura, codigo, nombre = fila[0], normalizar_codigo(fila[1]), fila[2]

        if abreviatura is None or codigo is None or nombre is None:
            continue

        abreviatura = str(abreviatura).strip()

        asignaturas[abreviatura] = Asignatura(
            codigo=codigo,
            abreviatura=abreviatura,
            nombre=str(nombre).strip(),
        )

    return asignaturas


def normalizar_codigo(valor):
    """
    Convierte el código numérico de la asignatura a texto.

    Devuelve None si el valor no es un código.
    """

    if valor is None:
        return None

    if isinstance(valor, int):
        return str(valor)

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    texto = str(valor).strip()

    if texto.isdigit():
        return texto

    return None
