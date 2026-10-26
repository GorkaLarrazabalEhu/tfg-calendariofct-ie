"""
Convierte los PDF de horarios semanales (una semana por página) a Excel.

Por cada página del PDF genera un bloque en la hoja "Horario":
  - Cabecera:  Kurtsoa/Curso, Taldea/Grupo, Astea/Semana
  - Horario:   franja horaria x días (Lunes..Viernes)
  - Asignaturas: acrónimo, código, nombre y profesorado

Por defecto lee los PDF de data/pdf/ y guarda los Excel en data/input/,
que es donde los busca el programa.

Uso (desde la raíz del proyecto):
    python scripts/pdf_a_excel.py                          # todos los PDF de data/pdf/
    python scripts/pdf_a_excel.py data/pdf/IE_3_Cast.pdf   # uno o varios PDF concretos
    python scripts/pdf_a_excel.py data/pdf/IE_3_Cast.pdf -o salida.xlsx

Requiere:  pip install pdfplumber openpyxl
"""

import argparse
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FINO = Side(style="thin")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
GRIS = PatternFill("solid", fgColor="DCDCDC")
CENTRADO = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQUIERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)
NEGRITA = Font(bold=True)

N_DIAS = 5  # columnas B..F del horario

RAIZ = Path(__file__).resolve().parents[1]
CARPETA_PDF = RAIZ / "data" / "pdf"
CARPETA_EXCEL = RAIZ / "data" / "input"


def limpiar(texto, unir_lineas=False):
    """Normaliza el texto de una celda del PDF."""
    if texto is None:
        return ""
    texto = texto.strip()
    if unir_lineas:
        texto = " ".join(texto.split())
    return texto


def quitar_columnas_vacias(tabla):
    """pdfplumber a veces detecta columnas fantasma, vacías en todas las filas."""
    ancho = max(len(fila) for fila in tabla)
    columnas = [
        i for i in range(ancho)
        if any(i < len(fila) and limpiar(fila[i]) for fila in tabla)
    ]
    return [[fila[i] if i < len(fila) else None for i in columnas] for fila in tabla]


def extraer_tablas(pagina):
    """
    Extrae las tablas de la página cortándola por la línea superior
    de la tabla de asignaturas.

    Si el horario es largo, queda pegado a la tabla de asignaturas y
    pdfplumber las une en una sola: las líneas verticales de una
    parten las columnas de la otra (p. ej. el lunes en tres columnas).
    """
    palabra = next((p for p in pagina.extract_words() if "ASIGNATURA" in p["text"]), None)
    if palabra is None:
        return pagina.extract_tables()

    y = max(
        (e["top"] for e in pagina.edges if e["orientation"] == "h" and e["top"] <= palabra["top"]),
        default=palabra["top"],
    ) - 0.5

    arriba = pagina.crop((0, 0, pagina.width, y))
    abajo = pagina.crop((0, y, pagina.width, pagina.height))
    return arriba.extract_tables() + abajo.extract_tables()


def normalizar_asignaturas(tabla):
    """
    Deja la tabla de asignaturas en cuatro columnas:
    acrónimo, código, asignatura y profesorado.

    Las celdas combinadas llegan como None y se descartan; si había
    columnas de más (tablas pegadas), así vuelven a quedar cuatro.
    """
    filas = [["", "", "ASIGNATURA/IRAKASGAIA", "PROFESOR/IRAKASLEA"]]
    for fila in tabla[1:]:
        valores = [c for c in fila if c is not None]
        if any(limpiar(c) for c in valores):
            filas.append(valores)
    return filas


def clasificar_tablas(tablas):
    """Identifica las tablas de cabecera, horario y asignaturas de una página."""
    cabecera = horario = asignaturas = None
    for t in tablas:
        primera = " ".join(c or "" for c in t[0])
        if "Kurtsoa" in primera or "Curso" in primera:
            cabecera = t
        elif "Lunes" in primera or "Astelehena" in primera:
            horario = t
        elif "ASIGNATURA" in primera or "IRAKASGAIA" in primera:
            asignaturas = normalizar_asignaturas(t)
    return cabecera, horario, asignaturas


def escribir_semana(ws, fila, cabecera, horario, asignaturas):
    """Escribe el bloque de una semana a partir de `fila`. Devuelve la siguiente fila libre."""
    ultima_col = 1 + N_DIAS

    # --- Cabecera (Curso / Grupo / Semana) ---
    if cabecera:
        celdas = [limpiar(c) for c in cabecera[0] if c]
        pares = [f"{celdas[i]} {celdas[i + 1]}" for i in range(0, len(celdas) - 1, 2)]
        texto = "     ".join(pares)
    else:
        texto = ""
    ws.cell(fila, 1, texto).font = Font(bold=True, size=11)
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima_col)
    fila += 1

    # --- Horario ---
    if horario:
        horario = quitar_columnas_vacias(horario)
        if len(horario[0]) != ultima_col:
            print(
                f"  [aviso] {texto}: la tabla tiene {len(horario[0])} columnas en vez de {ultima_col}",
                file=sys.stderr,
            )
        for i, row in enumerate(horario):
            row = (row + [""] * (ultima_col))[:ultima_col]
            for col, valor in enumerate(row, start=1):
                if i == 0:  # nombres de los días
                    valor = limpiar(valor, unir_lineas=True)
                else:
                    valor = limpiar(valor)
                c = ws.cell(fila, col, valor or None)
                c.border = BORDE
                c.alignment = CENTRADO
                if i == 0 or col == 1:
                    c.fill = GRIS
                    c.font = NEGRITA
            fila += 1
    fila += 1  # fila en blanco

    # --- Asignaturas ---
    if asignaturas:
        for i, row in enumerate(asignaturas):
            row = [limpiar(c, unir_lineas=True) for c in row] + [""] * 4
            acr, codigo, nombre, profesor = row[:4]
            if i > 0 and codigo.isdigit():
                codigo = int(codigo)
            valores = [acr, codigo, nombre, profesor]
            # A: acrónimo, B: código, C-D: asignatura, E-F: profesorado
            for col, valor in zip((1, 2, 3, 5), valores):
                c = ws.cell(fila, col, valor if valor != "" else None)
                c.alignment = IZQUIERDA if col in (3, 5) else CENTRADO
            ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
            ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=6)
            for col in range(1, ultima_col + 1):
                c = ws.cell(fila, col)
                c.border = BORDE
                if i == 0:
                    c.fill = GRIS
                    c.font = NEGRITA
            fila += 1

    return fila + 1  # fila en blanco entre semanas


def convertir(pdf_path, xlsx_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Horario"

    fila = 1
    semanas = 0
    with pdfplumber.open(pdf_path) as pdf:
        for n, pagina in enumerate(pdf.pages, start=1):
            cabecera, horario, asignaturas = clasificar_tablas(extraer_tablas(pagina))
            if not (cabecera or horario or asignaturas):
                print(f"  [aviso] página {n}: no se han encontrado tablas", file=sys.stderr)
                continue
            fila = escribir_semana(ws, fila, cabecera, horario, asignaturas)
            semanas += 1

    anchos = {"A": 10, "B": 22, "C": 22, "D": 22, "E": 22, "F": 22}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.sheet_view.zoomScale = 90

    wb.save(xlsx_path)
    print(f"{pdf_path.name} -> {xlsx_path.name}  ({semanas} semanas)")


def main():
    parser = argparse.ArgumentParser(description="Convierte PDF de horarios a Excel.")
    parser.add_argument("pdfs", nargs="*", type=Path, help=f"PDF a convertir (por defecto, todos los de {CARPETA_PDF})")
    parser.add_argument("-o", "--salida", type=Path, help=f"Excel de salida (solo con un PDF; por defecto, en {CARPETA_EXCEL})")
    args = parser.parse_args()

    pdfs = args.pdfs or sorted(CARPETA_PDF.glob("*.pdf"))
    if not pdfs:
        sys.exit(f"No se han encontrado PDF en {CARPETA_PDF}.")
    if args.salida and len(pdfs) > 1:
        sys.exit("-o/--salida solo se puede usar con un único PDF.")

    CARPETA_EXCEL.mkdir(parents=True, exist_ok=True)

    for pdf in pdfs:
        salida = args.salida or CARPETA_EXCEL / pdf.with_suffix(".xlsx").name
        convertir(pdf, salida)


if __name__ == "__main__":
    main()
